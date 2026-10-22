import base64
import copy
import io
import os
import re
import sqlite3
import traceback
from datetime import datetime
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd
import streamlit as st
import streamlit.components.v1 as components
import gspread
from google.oauth2.service_account import Credentials

st.set_page_config(page_title="샐러드 생산일보 및 원료코드 System", layout="wide")

DB_FILE = "stock_master.db"

# ---------------------------------------------------------
# 구글 시트 연동 함수 추가
# ---------------------------------------------------------
@st.cache_resource
def init_gspread():
    scopes = [
        "https://www.googleapis.com/auth/spreadsheets",
        "https://www.googleapis.com/auth/drive"
    ]
    
    gcp_secrets = st.secrets["gcp_service_account"]
    private_key = gcp_secrets.get("private_key", "")
    if "\\n" in private_key:
        private_key = private_key.replace("\\n", "\n")
    
    creds_dict = {
        "type": "service_account",
        "project_id": gcp_secrets.get("project_id", ""),
        "private_key_id": gcp_secrets.get("private_key_id", ""),
        "private_key": private_key,
        "client_email": gcp_secrets.get("client_email", ""),
        "client_id": gcp_secrets.get("client_id", ""),
        "auth_uri": "https://accounts.google.com/o/oauth2/auth",
        "token_uri": "https://oauth2.googleapis.com/token",
        "auth_provider_x509_cert_url": "https://www.googleapis.com/oauth2/v1/certs",
        "client_x509_cert_url": gcp_secrets.get("client_x509_cert_url", "")
    }

    creds = Credentials.from_service_account_info(creds_dict, scopes=scopes)
    client = gspread.authorize(creds)
    
    # 두 번째 수불부 앱 구글 시트 주소
    url = "https://docs.google.com/spreadsheets/d/1WM472gdj_7HBUvJdLRgwIU0t7GAMnZ47GpsCJfB5pBo/edit"
    return client.open_by_url(url)

# ---------------------------------------------------------
# DB 관리 함수
# ---------------------------------------------------------
def init_db():
    conn = sqlite3.connect(DB_FILE)
    cursor = conn.cursor()
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS closed_lots (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            item_name TEXT,
            yymmdd TEXT,
            avail_g REAL,
            vendor TEXT,
            closed_ym TEXT
        )
    """)
    conn.commit()
    conn.close()

def clear_db():
    conn = sqlite3.connect(DB_FILE)
    cursor = conn.cursor()
    cursor.execute("DELETE FROM closed_lots")
    conn.commit()
    conn.close()

def get_latest_closed_lots():
    conn = sqlite3.connect(DB_FILE)
    cursor = conn.cursor()
    cursor.execute("SELECT MAX(closed_ym) FROM closed_lots")
    max_ym = cursor.fetchone()[0]
    
    lots = {"양상추": [], "양배추": [], "적채": [], "프릴아이스": []}
    if max_ym:
        cursor.execute(
            "SELECT item_name, yymmdd, avail_g, vendor FROM closed_lots WHERE closed_ym = ?",
            (max_ym,)
        )
        rows = cursor.fetchall()
        for item_name, yymmdd, avail_g, vendor in rows:
            if item_name in lots:
                lots[item_name].append({
                    "yymmdd": yymmdd,
                    "avail_g": avail_g,
                    "avail": avail_g,
                    "vendor": vendor,
                    "label": f"{yymmdd[2:4]}/{yymmdd[4:]} 이월",
                    "is_from_db": True
                })
    conn.close()
    return max_ym, lots

def save_closed_lots(closed_ym, item_lots):
    conn = sqlite3.connect(DB_FILE)
    cursor = conn.cursor()
    cursor.execute("DELETE FROM closed_lots WHERE closed_ym = ?", (closed_ym,))
    
    for item_name, lot_list in item_lots.items():
        for lot in lot_list:
            avail_val = lot.get("avail", lot.get("avail_g", 0.0))
            if avail_val > 0:
                cursor.execute(
                    "INSERT INTO closed_lots (item_name, yymmdd, avail_g, vendor, closed_ym) VALUES (?, ?, ?, ?, ?)",
                    (item_name, lot["yymmdd"], avail_val, lot.get("vendor", "K01"), closed_ym)
                )
    conn.commit()
    conn.close()

init_db()

# ---------------------------------------------------------
# 원클릭 즉시 인쇄/PDF 저장 버튼 생성 함수
# ---------------------------------------------------------
def render_direct_print_button(html_payload, button_label="🖨️ 바로 출력 / PDF"):
    b64 = base64.b64encode(html_payload.encode('utf-8')).decode('utf-8')
    button_html = f"""
    <!DOCTYPE html>
    <html>
    <head>
    <style>
        body {{ margin: 0; padding: 0; font-family: "Source Sans Pro", sans-serif; overflow: hidden; }}
        .btn {{
            display: inline-flex; align-items: center; justify-content: center;
            font-weight: 400; padding: 0.25rem 0.75rem; border-radius: 0.5rem;
            min-height: 38.4px; margin: 0px; line-height: 1.6;
            width: 100%; cursor: pointer; font-size: 16px; box-sizing: border-box;
            transition: all 0.2s ease;
        }}
        @media (prefers-color-scheme: dark) {{
            body {{ background-color: transparent; }}
            .btn {{ background-color: #262730; color: #FAFAFA; border: 1px solid rgba(250, 250, 250, 0.2); }}
            .btn:hover {{ border-color: #FF4B4B; color: #FF4B4B; }}
        }}
        @media (prefers-color-scheme: light) {{
            body {{ background-color: transparent; }}
            .btn {{ background-color: #FFFFFF; color: #31333F; border: 1px solid rgba(49, 51, 63, 0.2); }}
            .btn:hover {{ border-color: #FF4B4B; color: #FF4B4B; }}
        }}
    </style>
    <script>
        function directPrint() {{
            try {{
                var b64 = "{b64}";
                var decoded = decodeURIComponent(escape(window.atob(b64)));
                
                var iframe = document.createElement('iframe');
                iframe.style.position = 'absolute';
                iframe.style.width = '0px';
                iframe.style.height = '0px';
                iframe.style.border = 'none';
                document.body.appendChild(iframe);
                
                var doc = iframe.contentWindow.document;
                doc.open();
                doc.write(decoded);
                doc.close();
                
                iframe.contentWindow.focus();
                
                setTimeout(function() {{ 
                    iframe.contentWindow.print(); 
                }}, 500);
                
            }} catch (e) {{
                alert("인쇄 화면을 여는 중 오류가 발생했습니다.");
            }}
        }}
    </script>
    </head>
    <body>
        <button class="btn" onclick="directPrint()">{button_label}</button>
    </body>
    </html>
    """
    components.html(button_html, height=45)

# ---------------------------------------------------------
# 일반 문서용 HTML 렌더링 함수
# ---------------------------------------------------------
def render_printable_document(title, df_data):
    html_table = df_data.to_html(index=False, classes="report-table")
    full_html = f"""
    <!DOCTYPE html>
    <html>
    <head>
    <meta charset="utf-8">
    <title>{title}</title>
    <style>
        body {{ font-family: '맑은 고딕', sans-serif; padding: 20px; color: #000; background-color: #fff; }}
        .header-container {{ display: flex; justify-content: space-between; align-items: center; margin-bottom: 20px; }}
        .report-title {{ font-size: 20pt; font-weight: bold; color: #1F4E79; margin: 0; }}
        
        .approval-table {{ border-collapse: collapse; text-align: center; font-size: 9pt; }}
        .approval-table th, .approval-table td {{ border: 1px solid #333; padding: 4px 8px; }}
        .approval-table th {{ background-color: #f2f2f2; width: 50px; }}
        .approval-table td {{ height: 40px; width: 60px; }}
        
        .report-table {{ width: 100%; border-collapse: collapse; font-size: 9pt; margin-top: 10px; }}
        .report-table th, .report-table td {{ border: 1px solid #333; padding: 6px 4px; text-align: center; }}
        .report-table th {{ background-color: #D9E1F2; font-weight: bold; color: #1F4E79; }}
        .report-table tr:nth-child(even) {{ background-color: #F9FAFB; }}
        
        @media print {{
            body {{ padding: 0; margin: 0; }}
            @page {{ size: A4 portrait; margin: 15mm; }}
        }}
    </style>
    </head>
    <body>
        <div class="header-container">
            <h1 class="report-title">{title}</h1>
            <table class="approval-table">
                <tr><th rowspan="2">결<br>재</th><th>담당</th><th>검토</th><th>승인</th></tr>
                <tr><td></td><td></td><td></td></tr>
            </table>
        </div>
        {html_table}
    </body>
    </html>
    """
    return full_html

# ---------------------------------------------------------
# 전처리 공정일지 전용 HTML 렌더링 함수
# ---------------------------------------------------------
def render_process_log_html(date_str, df_target, is_completed=False):
    try:
        dt = datetime.strptime(date_str, "%Y-%m-%d")
        yyyy, mm, dd = dt.strftime("%Y"), dt.strftime("%m"), dt.strftime("%d")
    except:
        yyyy, mm, dd = "202 ", "  ", "  "
    
    row_count = len(df_target)
    rows_html = ""
    
    if row_count == 0:
        rows_html = "<tr><td colspan='9' style='text-align:center; padding:15px;'>해당 일자의 데이터가 없습니다.</td></tr>"
    else:
        current_product = None
        for i, row in enumerate(df_target.to_dict('records')):
            item = row.get("품목", "")
            prod_name = row.get("제품명", "")
            
            # 절단 규격 분류
            if prod_name == "크라이 컷레터스":
                cut_spec = "절단 6X6cm"
            elif "양상추" in item or "프릴" in item:
                cut_spec = "절단 4X4cm"
            elif "양배추" in item or "적채" in item:
                cut_spec = "슬라이스 4mm"
            else:
                cut_spec = ""

            # 원료코드 분류
            if "프릴" in item:
                display_code = row.get("매입자료 기준 코드", "-")
            else:
                display_code = row.get("수불부 기준 코드", "-")

            # 다중 제품 rowspan 동적 그룹핑
            if prod_name != current_product:
                current_product = prod_name
                prod_row_count = len(df_target[df_target['제품명'] == prod_name])
                prod_cell = f'<td rowspan="{prod_row_count}"><b>{prod_name}</b></td>'
                check_cell = f'<td rowspan="{prod_row_count}" style="font-size:8pt; text-align:left; line-height:1.5;">부패,변질,잔잎 등의 비가식부위 제거<br>원재료 상태 및 이물제거 확인</td>'
            else:
                prod_cell = ""
                check_cell = ""

            # 투입량(후) 값 처리 - 완성본이 아닐 때는 단위 'kg'만 우측에 출력
            if is_completed:
                input_after_td = f'<td style="text-align:right;">{row.get("준비양(후)", 0)} kg</td>'
            else:
                input_after_td = '<td style="text-align:right;">kg</td>'

            rows_html += f"""
            <tr>
                {prod_cell}
                <td>{item}</td>
                <td>{display_code}</td>
                <td>적 / 부</td>
                {check_cell}
                <td>{cut_spec}</td>
                <td style="text-align:right;">{row.get("실투입량(전)", 0)} kg</td>
                {input_after_td}
                <td></td>
            </tr>
            """
    
    # 남은 여백 줄 추가
    for _ in range(5):
        rows_html += "<tr><td></td><td></td><td></td><td></td><td></td><td></td><td></td><td></td><td></td></tr>"

    html = f"""
    <!DOCTYPE html>
    <html>
    <head>
    <meta charset="utf-8">
    <title>전처리 공정일지</title>
    <style>
        body {{ font-family: '맑은 고딕', sans-serif; padding: 20px; color: #000; background-color: #fff; margin: 0 auto; max-width: 1200px; }}
        @media print {{
            body {{ padding: 0; margin: 0; max-width: none; }}
            @page {{ size: A4 landscape; margin: 15mm; }}
        }}
        .header-table {{ width: 100%; margin-bottom: 10px; }}
        .header-table td {{ vertical-align: top; }}
        .title {{ font-size: 26pt; font-weight: bold; letter-spacing: 2px; }}
        .approval-table {{ border-collapse: collapse; text-align: center; font-size: 10pt; float: right; margin-bottom: 5px; }}
        .approval-table th, .approval-table td {{ border: 1px solid #000; padding: 4px 8px; }}
        .approval-table th {{ width: 50px; font-weight: normal; }}
        .approval-table td {{ height: 40px; width: 60px; }}
        
        .main-table {{ width: 100%; border-collapse: collapse; font-size: 9pt; border: 1px solid #000; text-align: center; }}
        .main-table th, .main-table td {{ border: 1px solid #000; padding: 6px; }}
        .bg-gray {{ background-color: #f2f2f2; font-weight: bold; }}
    </style>
    </head>
    <body>
        <table class="header-table">
            <tr>
                <td class="title">전처리 공정일지</td>
                <td>
                    <table class="approval-table">
                        <tr><th rowspan="2">결<br>재</th><th>작 성</th><th>검 토</th><th>승 인</th></tr>
                        <tr><td></td><td></td><td></td></tr>
                    </table>
                </td>
            </tr>
        </table>

        <table class="main-table">
            <tr class="bg-gray">
                <td colspan="2">작성일자</td>
                <td colspan="2" style="background-color: #fff;">{yyyy} 년 &nbsp;&nbsp;&nbsp; {mm} 월 &nbsp;&nbsp;&nbsp; {dd} 일</td>
                <td colspan="2" style="background-color: #fff;">&nbsp;&nbsp; : &nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp; ~ &nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp; : &nbsp;&nbsp;</td>
                <td>작성자</td>
                <td colspan="2" style="background-color: #fff;"></td>
            </tr>
            <tr>
                <td colspan="2" rowspan="2" class="bg-gray">원료코드<br>부여방법</td>
                <td colspan="2" rowspan="2">원료 구분+입고일자 - 공급업체 구분</td>
                <td colspan="5" style="text-align: left; font-size: 8pt; padding: 4px 8px;">
                    카이피라(C), 프릴아이스(F), 버터헤드(B), 레드오크(O), 로메인(R), 케일(K), 적근대(S)<br>
                    치커리(E), 적채(A), 양배추(D), 양상추(X)
                </td>
            </tr>
            <tr>
                <td colspan="5" style="text-align: left; font-size: 8pt; padding: 4px 8px; border-top:1px dashed #999;">
                    공급업체 구분: 구름 (G01) , 에상스팜(H02), 승승장구(S01), 한스(K01)
                </td>
            </tr>
            <tr class="bg-gray">
                <th rowspan="2">제품명</th>
                <th rowspan="2">원료명</th>
                <th colspan="7">작업원물 정보</th>
            </tr>
            <tr class="bg-gray">
                <th>원료코드</th>
                <th>상태</th>
                <th>점검 항목</th>
                <th>절단규격</th>
                <th>생산량(전)</th>
                <th>투입량(후)</th>
                <th>비고</th>
            </tr>
            {rows_html}
        </table>
    </body>
    </html>
    """
    return html

st.title("🥗 샐러드 생산일보 원료코드 대조 및 수불부 연동 System")

# ---------------------------------------------------------
# 사이드바
# ---------------------------------------------------------
with st.sidebar:
    st.header("🗄️ 마스터 수불 DB 관리")
    last_ym, db_lots = get_latest_closed_lots()
    
    if last_ym:
        st.success(f"📌 현재 로드된 마감월: **{last_ym}**")
        with st.expander("이월 잔여 롯트 상세"):
            for k, v in db_lots.items():
                if v:
                    st.caption(f"**{k}**")
                    for l in v:
                        st.write(f"- 입고일: {l['yymmdd']} | 잔여: {round(l['avail_g']/1000.0, 1)}kg")
    else:
        st.info("💡 등록된 전월 마감 DB가 없습니다.\n당월 업로드 엑셀의 입고분부터 순서대로 소진됩니다.")

    st.markdown("---")
    if st.button("🗑️ 과거 이월 DB 완전 초기화", type="secondary"):
        clear_db()
        st.success("✅ DB가 초기화되었습니다!")
        st.rerun()

col1, col2 = st.columns(2)
with col1:
    excel_file = st.file_uploader("1. 샐러드 생산일지 (Excel) 업로드", type=["xlsx", "xls"])
with col2:
    csv_file = st.file_uploader("2. 원료 매입 입고자료 / 수불부 (CSV/Excel) 업로드", type=["csv", "xlsx", "xls"])

prefix_map = {"양상추": "X", "양배추": "D", "적채": "A", "프릴": "F", "프릴아이스": "F"}
vendor_code_map = {"한스": "K01", "승승장구": "S01", "에상스팜": "H02", "넥스토팜": "N01"}

# ---------------------------------------------------------
# 날짜 정밀 정제 함수
# ---------------------------------------------------------
def extract_yymmdd_and_mmdd(val):
    if pd.isna(val):
        today = datetime.now()
        return today.strftime("%y"), today.strftime("%y%m%d"), today.strftime("%m%d")
    if isinstance(val, (pd.Timestamp, datetime)):
        return val.strftime("%y"), val.strftime("%y%m%d"), val.strftime("%m%d")

    val_str = str(val).strip()
    
    m_dot_digit = re.search(r"(\d{2})\.(\d{2})(\d{2})", val_str)
    if m_dot_digit:
        yy, mm, dd = m_dot_digit.group(1), m_dot_digit.group(2), m_dot_digit.group(3)
        return yy, f"{yy}{mm}{dd}", f"{mm}{dd}"

    m_full = re.search(r"(\d{2,4})[-./](\d{1,2})[-./](\d{1,2})", val_str)
    if m_full:
        p1 = m_full.group(1)
        yy = p1[2:] if len(p1) == 4 else p1.zfill(2)
        mm = m_full.group(2).zfill(2)
        dd = m_full.group(3).zfill(2)
        return yy, f"{yy}{mm}{dd}", f"{mm}{dd}"

    m_digits = re.search(r"\b(\d{6})\b", val_str)
    if m_digits:
        d = m_digits.group(1)
        return d[:2], d, d[2:]

    current_yy = datetime.now().strftime("%y")
    m_short = re.search(r"(\d{1,2})\s*[-./월]\s*(\d{1,2})", val_str)
    if m_short:
        mm = m_short.group(1).zfill(2)
        dd = m_short.group(2).zfill(2)
        return current_yy, f"{current_yy}{mm}{dd}", f"{mm}{dd}"

    try:
        dt = pd.to_datetime(val_str, errors="coerce", yearfirst=True)
        if pd.notna(dt):
            return dt.strftime("%y"), dt.strftime("%y%m%d"), dt.strftime("%m%d")
    except Exception:
        pass

    today = datetime.now()
    return current_yy, today.strftime("%y%m%d"), today.strftime("%m%d")

def find_stock_sheet_name(xls):
    for sheet_name in xls.sheet_names:
        if sheet_name.strip().lower() == "sheet1": return sheet_name
    keywords = ["양상추", "양배추", "적채"]
    for sheet_name in xls.sheet_names:
        try:
            df = pd.read_excel(xls, sheet_name=sheet_name, header=None, nrows=10)
            for idx, row in df.iterrows():
                row_str = " ".join(row.dropna().astype(str).tolist())
                if all(k in row_str for k in keywords):
                    return sheet_name
        except Exception: continue
    return xls.sheet_names[0]

def parse_and_convert_to_standard_ledger(excel_file):
    xls = pd.ExcelFile(excel_file)
    stock_sheet = find_stock_sheet_name(xls)
    df_raw = pd.read_excel(xls, sheet_name=stock_sheet, header=None)

    def parse_cell(val):
        if pd.isna(val): return 0.0
        try:
            c = re.sub(r"[^\d.]", "", str(val))
            return float(c) if c else 0.0
        except Exception: return 0.0

    raw_parsed_list = []
    for idx, row in df_raw.iterrows():
        if len(row) == 0 or pd.isna(row.iloc[0]): continue
        a_val = str(row.iloc[0]).strip()
        if not a_val or "날짜" in a_val or "입고" in a_val or "매입" in a_val: continue

        _, yymmdd, _ = extract_yymmdd_and_mmdd(a_val)
        formatted_date = f"20{yymmdd[:2]}-{yymmdd[2:4]}-{yymmdd[4:]}"

        y_in  = parse_cell(row.iloc[1]) if len(row) > 1 else 0.0
        y_stk = parse_cell(row.iloc[2]) if len(row) > 2 else 0.0
        y_use = parse_cell(row.iloc[3]) if len(row) > 3 else 0.0

        b_in  = parse_cell(row.iloc[4]) if len(row) > 4 else 0.0
        b_stk = parse_cell(row.iloc[5]) if len(row) > 5 else 0.0
        b_use = parse_cell(row.iloc[6]) if len(row) > 6 else 0.0

        j_in  = parse_cell(row.iloc[7]) if len(row) > 7 else 0.0
        j_stk = parse_cell(row.iloc[8]) if len(row) > 8 else 0.0
        j_use = parse_cell(row.iloc[9]) if len(row) > 9 else 0.0

        raw_parsed_list.append({
            "date": formatted_date, "yymmdd": yymmdd,
            "y_in": y_in, "y_use": y_use, "y_stk": y_stk,
            "b_in": b_in, "b_use": b_use, "b_stk": b_stk,
            "j_in": j_in, "j_use": j_use, "j_stk": j_stk
        })

    grouped_by_date = {}
    date_order = []

    for item in raw_parsed_list:
        dt = item["date"]
        if dt not in grouped_by_date:
            grouped_by_date[dt] = {
                "date": dt, "yymmdd": item["yymmdd"],
                "y_in": 0.0, "y_use": 0.0, "y_stk": 0.0,
                "b_in": 0.0, "b_use": 0.0, "b_stk": 0.0,
                "j_in": 0.0, "j_use": 0.0, "j_stk": 0.0
            }
            date_order.append(dt)

        grouped_by_date[dt]["y_in"] += item["y_in"]
        grouped_by_date[dt]["b_in"] += item["b_in"]
        grouped_by_date[dt]["j_in"] += item["j_in"]

        grouped_by_date[dt]["y_use"] += item["y_use"]
        grouped_by_date[dt]["b_use"] += item["b_use"]
        grouped_by_date[dt]["j_use"] += item["j_use"]

        if item["y_stk"] > 0: grouped_by_date[dt]["y_stk"] = item["y_stk"]
        if item["b_stk"] > 0: grouped_by_date[dt]["b_stk"] = item["b_stk"]
        if item["j_stk"] > 0: grouped_by_date[dt]["j_stk"] = item["j_stk"]

    _, db_lots = get_latest_closed_lots()
    item_queues = {
        "yangsangchu": copy.deepcopy(db_lots["양상추"]),
        "yangbaechu": copy.deepcopy(db_lots["양배추"]),
        "jeokchae": copy.deepcopy(db_lots["적채"])
    }

    data_rows = []
    generated_code_rows = []
    target_items_order = ["양상추", "양배추", "적채"]

    for dt in date_order:
        g = grouped_by_date[dt]
        yymmdd = g["yymmdd"]
        mmdd_str = f"{yymmdd[2:4]}/{yymmdd[4:]}"

        if g["y_in"] > 0: item_queues["yangsangchu"].append({"yymmdd": yymmdd, "label": f"{mmdd_str} 매입", "avail": g["y_in"], "vendor": "K01"})
        if g["b_in"] > 0: item_queues["yangbaechu"].append({"yymmdd": yymmdd, "label": f"{mmdd_str} 매입", "avail": g["b_in"], "vendor": "K01"})
        if g["j_in"] > 0: item_queues["jeokchae"].append({"yymmdd": yymmdd, "label": f"{mmdd_str} 매입", "avail": g["j_in"], "vendor": "K01"})

        def trace_usage(item_key, item_prefix, use_amount):
            if use_amount <= 0: return "-", "-"
            rem = use_amount
            used_sources, used_codes = [], []

            for lot in item_queues[item_key]:
                if rem <= 0: break
                avail_val = lot.get("avail", 0.0)
                if avail_val <= 0: continue
                
                u = min(avail_val, rem)
                lot["avail"] = avail_val - u
                rem -= u

                vendor = lot.get("vendor", "K01")
                used_sources.append(f"{lot['label']}({round(u/1000.0, 1)}kg)")
                used_codes.append(f"{item_prefix}{lot['yymmdd']}-{vendor}")

            if rem > 0:
                used_sources.append(f"부족({round(rem/1000.0, 1)}kg)")
                used_codes.append(f"{item_prefix}{yymmdd}-K01[부족]")

            return " / ".join(used_sources), " / ".join(list(dict.fromkeys(used_codes)))

        y_src, y_code = trace_usage("yangsangchu", "X", g["y_use"])
        b_src, b_code = trace_usage("yangbaechu", "D", g["b_use"])
        j_src, j_code = trace_usage("jeokchae", "A", g["j_use"])

        note = "휴무" if (g["y_in"] == 0 and g["y_use"] == 0 and g["b_in"] == 0 and g["b_use"] == 0 and g["j_in"] == 0 and g["j_use"] == 0) else ""

        row_data = {
            "date": dt, "yymmdd": yymmdd,
            "yangsangchu_in": round(g["y_in"] / 1000.0, 1), 
            "yangsangchu_use": round(g["y_use"] / 1000.0, 1), 
            "yangsangchu_source": y_src, 
            "yangsangchu_code": y_code,
            "yangsangchu_stock": round(g["y_stk"] / 1000.0, 1),
            "yangbaechu_in": round(g["b_in"] / 1000.0, 1), 
            "yangbaechu_use": round(g["b_use"] / 1000.0, 1), 
            "yangbaechu_source": b_src, 
            "yangbaechu_code": b_code,
            "yangbaechu_stock": round(g["b_stk"] / 1000.0, 1),
            "jeokchae_in": round(g["j_in"] / 1000.0, 1), 
            "jeokchae_use": round(g["j_use"] / 1000.0, 1), 
            "jeokchae_source": j_src, 
            "jeokchae_code": j_code,
            "jeokchae_stock": round(g["j_stk"] / 1000.0, 1),
            "note": note
        }
        data_rows.append(row_data)

        code_map = {"양상추": y_code, "양배추": b_code, "적채": j_code}
        src_map = {"양상추": y_src, "양배추": b_src, "적채": j_src}
        use_map = {"양상추": g["y_use"], "양배추": g["b_use"], "적채": g["j_use"]}

        for item_name in target_items_order:
            u_val = use_map[item_name]
            if u_val > 0:
                generated_code_rows.append({
                    "날짜": dt, "원재료": item_name, "생성코드": code_map.get(item_name, "-"),
                    "투입량(kg)": round(u_val / 1000.0, 1), "준비량(kg)": 0.0, "소진매입일(상세)": src_map.get(item_name, "-"), "비고": note
                })

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "원재료 수불부"
    ws.views.sheetView[0].showGridLines = True
    ws.merge_cells("A1:Q1")
    ws["A1"] = "원재료 수불부 (Raw Material Inventory Ledger - kg 단위)"
    ws["A1"].font = Font(name="맑은 고딕", size=16, bold=True, color="1F4E79")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    header_fill_main = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_fill_sub = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    regular_font = Font(name="맑은 고딕", size=10)
    thin_border = Border(left=Side(style="thin", color="D9D9D9"), right=Side(style="thin", color="D9D9D9"),
                         top=Side(style="thin", color="D9D9D9"), bottom=Side(style="thin", color="D9D9D9"))

    headers_main = [("A2:A3", "일자"), ("B2:F2", "양상추(kg)"), ("G2:K2", "양배추(kg)"), ("L2:P2", "적채(kg)"), ("Q2:Q3", "비고")]
    for cell_range, title in headers_main:
        ws.merge_cells(cell_range)
        ws[cell_range.split(":")[0]] = title

    for col_idx in range(1, 18):
        c = ws.cell(row=2, column=col_idx)
        c.fill, c.font, c.alignment = header_fill_main, Font(name="맑은 고딕", size=11, bold=True, color="FFFFFF"), Alignment(horizontal="center", vertical="center")

    sub_headers = ["기초재고", "입고량", "투입량", "소진 매입일자(상세)", "기말재고"]
    for start_col in [2, 7, 12]:
        for idx, sub in enumerate(sub_headers):
            c = ws.cell(row=3, column=start_col + idx, value=sub)
            c.fill, c.font, c.alignment = header_fill_sub, Font(name="맑은 고딕", size=10, bold=True, color="1F4E79"), Alignment(horizontal="center", vertical="center")

    start_row = 4
    for r_idx, item in enumerate(data_rows, start=start_row):
        ws.cell(row=r_idx, column=1, value=item["date"]).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=r_idx, column=17, value=item.get("note", "")).alignment = Alignment(horizontal="left", vertical="center")

        items_map = [
            (2, item.get("yangsangchu_in", 0), item.get("yangsangchu_use", 0), item.get("yangsangchu_source", "-"), item.get("yangsangchu_stock", 0)),
            (7, item.get("yangbaechu_in", 0), item.get("yangbaechu_use", 0), item.get("yangbaechu_source", "-"), item.get("yangbaechu_stock", 0)),
            (12, item.get("jeokchae_in", 0), item.get("jeokchae_use", 0), item.get("jeokchae_source", "-"), item.get("jeokchae_stock", 0))
        ]

        for start_col, in_val, use_val, source_str, stock_val in items_map:
            end_col_let = get_column_letter(start_col + 4)
            if r_idx == start_row:
                ws.cell(row=r_idx, column=start_col, value=round(stock_val + use_val - in_val, 1))
            else:
                ws.cell(row=r_idx, column=start_col, value=f"={end_col_let}{r_idx-1}")

            ws.cell(row=r_idx, column=start_col + 1, value=in_val)
            ws.cell(row=r_idx, column=start_col + 2, value=use_val)
            ws.cell(row=r_idx, column=start_col + 3, value=source_str).alignment = Alignment(horizontal="center", vertical="center")
            ws.cell(row=r_idx, column=start_col + 4, value=stock_val)

        for c_idx in range(1, 18):
            cell = ws.cell(row=r_idx, column=c_idx)
            cell.font, cell.border = regular_font, thin_border
            if c_idx in [2, 3, 4, 6, 7, 8, 9, 11, 12, 13, 14, 16]:
                cell.number_format = "#,##0.0"
                cell.alignment = Alignment(horizontal="right", vertical="center")

    buffer = io.BytesIO()
    wb.save(buffer)

    last_row_data = data_rows[-1] if data_rows else {}
    target_stocks = {
        "양상추": last_row_data.get("yangsangchu_stock", 0.0) * 1000.0,
        "양배추": last_row_data.get("yangbaechu_stock", 0.0) * 1000.0,
        "적채": last_row_data.get("jeokchae_stock", 0.0) * 1000.0
    }

    final_remaining_lots = {"양상추": [], "양배추": [], "적채": []}

    for item_k, q_key in [("양상추", "yangsangchu"), ("양배추", "yangbaechu"), ("적채", "jeokchae")]:
        target_g = target_stocks[item_k]
        valid_q = [q for q in item_queues[q_key] if q.get("avail", 0.0) > 0]
        rem_target = target_g
        for q in reversed(valid_q):
            if rem_target <= 0: break
            assign_g = min(q["avail"], rem_target)
            final_remaining_lots[item_k].insert(0, {
                "yymmdd": q["yymmdd"],
                "avail_g": assign_g,
                "avail": assign_g,
                "vendor": q.get("vendor", "K01")
            })
            rem_target -= assign_g

    return buffer.getvalue(), data_rows, generated_code_rows, stock_sheet, final_remaining_lots

def parse_csv_purchase(uploaded_file):
    if uploaded_file is None: return pd.DataFrame()
    file_name = uploaded_file.name.lower()
    
    # 엑셀 파일인 경우 모든 시트를 병합하여 반환
    if file_name.endswith(".xlsx") or file_name.endswith(".xls"):
        try:
            df_dict = pd.read_excel(uploaded_file, sheet_name=None)
            return pd.concat(df_dict.values(), ignore_index=True)
        except Exception: return pd.DataFrame()

    encodings = ["utf-8-sig", "cp949", "euc-kr", "utf-8", "latin1"]
    for enc in encodings:
        try:
            uploaded_file.seek(0)
            return pd.read_csv(uploaded_file, encoding=enc)
        except Exception: continue
    return pd.DataFrame()

def generate_salad_report(excel_file, csv_file, code_summary_rows, stock_sheet_name):
    xls = pd.ExcelFile(excel_file)
    df_csv = parse_csv_purchase(csv_file)

    has_csv = not df_csv.empty

    purchase_records = []
    if has_csv:
        date_col = next((c for c in df_csv.columns if "기간" in str(c) or "일자" in str(c) or "날짜" in str(c)), None)
        item_col = next((c for c in df_csv.columns if "품명" in str(c) or "품목" in str(c) or "원재료" in str(c)), None)
        vendor_col = next((c for c in df_csv.columns if "거래처" in str(c) or "매입처" in str(c) or "공급처" in str(c)), None)

        if date_col and item_col:
            for _, c_row in df_csv.iterrows():
                p_date_raw = c_row[date_col]
                p_item_raw = str(c_row[item_col]).strip()
                p_vendor_raw = str(c_row[vendor_col]).strip() if vendor_col and pd.notna(c_row[vendor_col]) else ""

                if pd.isna(p_date_raw) or not p_item_raw: continue
                if p_vendor_raw == "기타": continue

                _, p_yymmdd, _ = extract_yymmdd_and_mmdd(p_date_raw)
                p_std_dt = f"20{p_yymmdd[:2]}-{p_yymmdd[2:4]}-{p_yymmdd[4:]}"

                if "양상추" in p_item_raw: p_item = "양상추"
                elif "양배추" in p_item_raw: p_item = "양배추"
                elif "적채" in p_item_raw: p_item = "적채"
                elif "프릴" in p_item_raw: p_item = "프릴아이스"
                else: p_item = p_item_raw

                pref = prefix_map.get(p_item, prefix_map.get("프릴" if "프릴" in p_item else "양상추", "X"))
                v_code = vendor_code_map.get(p_vendor_raw, "K01")
                p_code = f"{pref}{p_yymmdd}-{v_code}"

                purchase_records.append({
                    "std_date": p_std_dt,
                    "yymmdd": p_yymmdd,
                    "item": p_item,
                    "code": p_code,
                    "vendor": p_vendor_raw
                })

    code_summary_map = {}
    for r in code_summary_rows:
        raw_dt = r.get("날짜", "")
        _, yymmdd, _ = extract_yymmdd_and_mmdd(raw_dt)
        std_dt = f"20{yymmdd[:2]}-{yymmdd[2:4]}-{yymmdd[4:]}"
        item = r.get("원재료", "").strip()
        c_code = r.get("생성코드", "-")
        code_summary_map[(std_dt, item)] = c_code

    exclude_sheets = ["테스트", "원가", stock_sheet_name, "Sheet1", "sheet1"]
    daily_sheets = [
        s for s in xls.sheet_names 
        if s not in exclude_sheets and not s.startswith("재고") and not re.match(r"^\d{2}\.\d{2}$", s.strip())
    ]

    records = []
    for sheet in daily_sheets:
        try:
            df_s = pd.read_excel(excel_file, sheet_name=sheet, header=None)
            if df_s.shape[0] < 5 or df_s.shape[1] < 14: continue
            
            prod_date = df_s.iloc[1, 0]
            if pd.isna(prod_date) or "작업" in str(prod_date):
                prod_date = df_s.iloc[0, 0]

            raw_product_name = str(df_s.iloc[1, 1]).strip()
            normalized_name = raw_product_name.replace(" ", "").replace(",", "")
            if "샐러드믹스1000" in normalized_name:
                product_name = "브런치빈 샐러드믹스"
            else:
                product_name = raw_product_name if raw_product_name and raw_product_name.lower() != "nan" else "-"

            items = [df_s.iloc[1, 4], df_s.iloc[1, 7], df_s.iloc[1, 10], df_s.iloc[1, 13]]
            lots = [df_s.iloc[2, 5], df_s.iloc[2, 8], df_s.iloc[2, 11], df_s.iloc[2, 14]]
            prep_pck = [df_s.iloc[3, 4], df_s.iloc[3, 7], df_s.iloc[3, 10], df_s.iloc[3, 13]]
            actual_input = [df_s.iloc[5, 4], df_s.iloc[5, 7], df_s.iloc[5, 10], df_s.iloc[5, 13]]

            for item, lot, prep, act in zip(items, lots, prep_pck, actual_input):
                if pd.isna(item) or str(item).strip() == "" or str(item).strip().lower() == "nan": continue
                try: p_val = float(prep) if pd.notna(prep) else 0.0
                except (ValueError, TypeError): p_val = 0.0

                try: a_val = float(act) if pd.notna(act) else 0.0
                except (ValueError, TypeError): a_val = 0.0

                lot_str = str(lot).strip() if pd.notna(lot) else ""
                has_lot = (lot_str != "" and lot_str.lower() != "nan" and lot_str.lower() != "none")
                
                _, prod_yymmdd, _ = extract_yymmdd_and_mmdd(prod_date)
                formatted_prod_date = f"20{prod_yymmdd[:2]}-{prod_yymmdd[2:4]}-{prod_yymmdd[4:]}"

                records.append({
                    "생산일자": formatted_prod_date, 
                    "제품명": product_name,
                    "prod_yymmdd": prod_yymmdd, 
                    "품목": str(item).strip(),
                    "작업일지_롯트": lot_str if has_lot else "없음", "act_g": a_val,
                    "실투입량(전)": round(a_val / 1000.0, 1), "준비양(후)": round(p_val / 1000.0, 1),
                    "has_lot": has_lot,
                })
                
            if df_s.shape[0] > 7 and df_s.shape[1] > 4:
                raw_prod2 = str(df_s.iloc[7, 0]).strip()
                norm_prod2 = raw_prod2.replace(" ", "")
                if "컷레터스" in norm_prod2:
                    act2_raw = df_s.iloc[7, 4]
                    try:
                        act2_val = float(act2_raw) if pd.notna(act2_raw) else 0.0
                    except:
                        act2_val = 0.0
                    
                    prep2_raw = df_s.iloc[7, 1]
                    try:
                        prep2_val = float(prep2_raw) if pd.notna(prep2_raw) else 0.0
                    except:
                        prep2_val = 0.0
                        
                    if act2_val > 0 or prep2_val > 0:
                        records.append({
                            "생산일자": formatted_prod_date, 
                            "제품명": "크라이 컷레터스", 
                            "prod_yymmdd": prod_yymmdd, 
                            "품목": "양상추",
                            "작업일지_롯트": "수불부코드자동적용", 
                            "act_g": act2_val,
                            "실투입량(전)": round(act2_val / 1000.0, 1), 
                            "준비양(후)": round(prep2_val / 1000.0, 1),
                            "has_lot": True, 
                        })

        except Exception: continue

    if not records: return pd.DataFrame()
    records.sort(key=lambda x: x["prod_yymmdd"])

    final_rows = []
    for rec in records:
        raw_item, lot, has_lot = rec["품목"], rec["작업일지_롯트"], rec["has_lot"]
        prod_date, prod_yymmdd = rec["생산일자"], rec["prod_yymmdd"]
        
        is_frill = "프릴" in raw_item

        if "양상추" in raw_item: lookup_item = "양상추"
        elif "양배추" in raw_item: lookup_item = "양배추"
        elif "적채" in raw_item: lookup_item = "적채"
        elif is_frill: lookup_item = "프릴아이스"
        else: lookup_item = raw_item.strip()

        p = prefix_map.get(lookup_item, prefix_map.get("프릴" if is_frill else "양상추", "X"))

        code_subul = "-"
        if not is_frill:
            code_subul = code_summary_map.get((prod_date, lookup_item), "-")

        code_purchase = "-"
        if has_csv:
            valid_p = [pr for pr in purchase_records if pr["item"] == lookup_item and pr["std_date"] <= prod_date]
            if valid_p:
                valid_p.sort(key=lambda x: x["std_date"], reverse=True)
                code_purchase = valid_p[0]["code"]

        if rec["제품명"] == "크라이 컷레터스":
            code_worklog = code_subul
            is_match = "일치(수불부 자동부여)"
        else:
            if has_lot:
                _, yymmdd_worklog, _ = extract_yymmdd_and_mmdd(lot)
                code_worklog = f"{p}{yymmdd_worklog}-K01" if yymmdd_worklog else "없음"
            else: 
                code_worklog = "없음"

            if is_frill:
                if code_worklog != "없음" and code_purchase != "-" and (code_worklog in code_purchase or code_worklog.split("-")[0] in code_purchase):
                    is_match = "일치"
                elif code_worklog == "없음":
                    is_match = "불일치(작업일지 미기재)"
                else:
                    is_match = "불일치(수불부 대상 제외 품목)"
            else:
                if code_worklog != "없음" and code_subul != "-" and code_worklog in code_subul:
                    is_match = "일치"
                elif code_worklog == "없음":
                    is_match = "불일치(작업일지 미기재)"
                elif code_subul == "-":
                    is_match = "불일치(수불부 소진 미발생)"
                else:
                    is_match = "불일치(혼합/추측)"

        rec["작업일지 기준 코드"], rec["수불부 기준 코드"], rec["매입자료 기준 코드"], rec["일치여부"] = code_worklog, code_subul, code_purchase, is_match
        final_rows.append(rec)

    df_report = pd.DataFrame(final_rows)
    return df_report[["생산일자", "제품명", "품목", "작업일지 기준 코드", "수불부 기준 코드", "매입자료 기준 코드", "일치여부", "실투입량(전)", "준비양(후)"]]

# ---------------------------------------------------------
# Streamlit 화면 구성 (6개 탭)
# ---------------------------------------------------------
tab1, tab2, tab3, tab4, tab5, tab6 = st.tabs([
    "🔍 원료코드 대조표 (자동 연동)",
    "📊 표준 원재료 수불부 정리",
    "🏷️ 원료코드 생성 및 일자별 정리",
    "🔒 당월 수불 마감 확정",
    "📋 전처리 공정일지 출력",
    "📋 전처리 공정일지 (완성본)"
])

if excel_file is not None:
    try:
        excel_bytes, standard_rows, code_summary_rows, stock_sheet_name, final_lots = parse_and_convert_to_standard_ledger(excel_file)
        df_result = generate_salad_report(excel_file, csv_file, code_summary_rows, stock_sheet_name)

        # --- Tab 1: 원료코드 대조표 ---
        with tab1:
            if not df_result.empty:
                st.success("✅ 소진 추적이 연동된 표준 수불부 및 매입자료 기반 원료코드 대조가 완료되었습니다!")
                
                st.markdown("##### 📌 (단위: kg)")
                col_d, col_p = st.columns([3, 1])
                with col_d:
                    buffer = io.BytesIO()
                    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
                        df_result.to_excel(writer, index=False, sheet_name="원료코드_대조표")
                    st.download_button("📥 엑셀 다운로드 (.xlsx)", data=buffer.getvalue(), file_name="원료코드_대조표.xlsx", use_container_width=True)
                with col_p:
                    html_content = render_printable_document("샐러드 생산일보 원료코드 대조표 (단위: kg)", df_result)
                    render_direct_print_button(html_content, "🖨️ 바로 출력 / PDF 저장")
                
                st.dataframe(df_result, use_container_width=True)

        # --- Tab 2: 표준 원재료 수불부 ---
        with tab2:
            st.subheader("📋 표준 원재료 수불부 (kg 단위 적용)")
            df_preview = pd.DataFrame(standard_rows)
            if "yymmdd" in df_preview.columns:
                df_preview = df_preview.drop(columns=["yymmdd"])

            ordered_cols = [
                "date",
                "yangsangchu_in", "yangsangchu_use", "yangsangchu_source", "yangsangchu_stock",
                "yangbaechu_in", "yangbaechu_use", "yangbaechu_source", "yangbaechu_stock",
                "jeokchae_in", "jeokchae_use", "jeokchae_source", "jeokchae_stock",
                "note"
            ]
            df_preview = df_preview[[col for col in ordered_cols if col in df_preview.columns]]

            korean_column_map = {
                "date": "일자",
                "yangsangchu_in": "양상추_입고량(kg)", "yangsangchu_use": "양상추_투입량(kg)", "yangsangchu_source": "양상추_소진 매입일자(상세)", "yangsangchu_stock": "양상추_기말재고(kg)",
                "yangbaechu_in": "양배추_입고량(kg)", "yangbaechu_use": "양배추_투입량(kg)", "yangbaechu_source": "양배추_소진 매입일자(상세)", "yangbaechu_stock": "양배추_기말재고(kg)",
                "jeokchae_in": "적채_입고량(kg)", "jeokchae_use": "적채_투입량(kg)", "jeokchae_source": "적채_소진 매입일자(상세)", "jeokchae_stock": "적채_기말재고(kg)",
                "note": "비고",
            }
            df_preview = df_preview.rename(columns=korean_column_map)

            col_d2, col_p2 = st.columns([3, 1])
            with col_d2:
                st.download_button("📥 엑셀 다운로드 (.xlsx)", data=excel_bytes, file_name="표준_원재료_수불부(kg).xlsx", use_container_width=True)
            with col_p2:
                html_content = render_printable_document("표준 원재료 수불부 (단위: kg)", df_preview)
                render_direct_print_button(html_content, "🖨️ 바로 출력 / PDF 저장")
            
            st.dataframe(df_preview, use_container_width=True)

            # --- Tab 2 내부 가장 하단에 구글 시트 전송 버튼 추가 (출고량만 덮어쓰기 전송) ---
            st.markdown("---")
            st.subheader("☁️ 메인 수불부 구글 시트로 출고량 연동 (덮어쓰기 방식)")
            st.info("💡 엑셀에 포함된 일자의 기존 '시스템자동' 연동 데이터는 지워지고 최신 '출고(사용량)' 데이터로 덮어씌워집니다. (수기 입력 데이터는 보존됨)")
            
            if st.button("🚀 현재 엑셀 데이터로 구글 시트 출고량 업데이트하기", type="primary", use_container_width=True):
                with st.spinner("구글 시트의 기존 데이터를 정리하고 최신 출고 데이터를 전송 중입니다..."):
                    try:
                        doc = init_gspread()
                        sheet = doc.worksheet("시트1") # 수불부 시트 이름에 맞게 확인
                        
                        # 1. 업로드된 엑셀 데이터의 고유 날짜들 추출
                        excel_dates = set(str(row["date"]) for row in standard_rows)
                        
                        if not excel_dates:
                            st.warning("전송할 데이터가 없습니다.")
                        else:
                            # 2. 구글 시트 전체 데이터 가져오기
                            all_values = sheet.get_all_values()
                            
                            # 3. 삭제할 행 인덱스 찾기 (조건: 엑셀에 있는 날짜 AND 시스템자동 입력건)
                            rows_to_delete = []
                            for i, row_data in enumerate(all_values):
                                if i == 0: 
                                    continue
                                
                                sheet_date = row_data[0].strip() if len(row_data) > 0 else ""
                                sheet_vendor = row_data[2].strip() if len(row_data) > 2 else ""
                                
                                if sheet_date in excel_dates and sheet_vendor == "시스템자동":
                                    rows_to_delete.append(i + 1)
                            
                            # 4. 행 삭제 (아래쪽 행부터 지워야 인덱스가 꼬이지 않음)
                            for row_idx in sorted(rows_to_delete, reverse=True):
                                sheet.delete_row(row_idx)
                            
                            # 5. 새로운 '출고' 데이터 생성
                            rows_to_upload = []
                            for row in standard_rows:
                                date_str = str(row["date"])
                                
                                # [양상추] 출고량만 전송
                                if row.get("yangsangchu_use", 0) > 0:
                                    rows_to_upload.append([date_str, "출고", "시스템자동", "양상추", row["yangsangchu_use"], 0, 0, row.get("yangsangchu_code", "")])
                                    
                                # [양배추] 출고량만 전송
                                if row.get("yangbaechu_use", 0) > 0:
                                    rows_to_upload.append([date_str, "출고", "시스템자동", "양배추", row["yangbaechu_use"], 0, 0, row.get("yangbaechu_code", "")])
                                    
                                # [적채] 출고량만 전송
                                if row.get("jeokchae_use", 0) > 0:
                                    rows_to_upload.append([date_str, "출고", "시스템자동", "적채", row["jeokchae_use"], 0, 0, row.get("jeokchae_code", "")])

                            # 6. 새 데이터 구글 시트에 일괄 추가
                            if rows_to_upload:
                                sheet.append_rows(rows_to_upload)
                                st.success(f"✅ 기존 중복 출고 데이터 {len(rows_to_delete)}건 삭제 후, 최신 출고 데이터 {len(rows_to_upload)}건이 성공적으로 업데이트되었습니다!")
                                st.balloons()
                            else:
                                if len(rows_to_delete) > 0:
                                     st.success(f"✅ 기존 중복 출고 데이터 {len(rows_to_delete)}건 삭제 완료. (새로 전송할 출고 데이터는 없습니다.)")
                                else:
                                     st.warning("전송할 출고 수량(0 초과) 데이터가 없습니다.")
                            
                    except Exception as e:
                        st.error(f"구글 시트 연동 중 오류가 발생했습니다: {e}")

        # --- Tab 3: 원료코드 생성 및 일자별 정리 ---
        with tab3:
            st.subheader("🏷️ 소진 매입재고 기반 일자별 원료코드 생성 결과")
            df_code_summary = pd.DataFrame(code_summary_rows)
            if not df_code_summary.empty:
                column_order = ["날짜", "원재료", "생성코드", "투입량(kg)", "준비량(kg)", "소진매입일(상세)", "비고"]
                df_code_summary = df_code_summary[[c for c in column_order if c in df_code_summary.columns]]

                item_priority = {"양상추": 1, "양배추": 2, "적채": 3}
                df_code_summary["_sort_key"] = df_code_summary["원재료"].map(lambda x: item_priority.get(x, 99))
                df_code_summary = df_code_summary.sort_values(by=["날짜", "_sort_key"]).drop(columns=["_sort_key"])

                col_d3, col_p3 = st.columns([3, 1])
                with col_d3:
                    code_buffer = io.BytesIO()
                    with pd.ExcelWriter(code_buffer, engine="openpyxl") as writer:
                        df_code_summary.to_excel(writer, index=False, sheet_name="일자별_원료코드_생성")
                    st.download_button("📥 엑셀 다운로드 (.xlsx)", data=code_buffer.getvalue(), file_name="일자별_생성_원료코드_목록.xlsx", use_container_width=True)
                with col_p3:
                    html_content = render_printable_document("일자별 생성 원료코드 목록", df_code_summary)
                    render_direct_print_button(html_content, "🖨️ 바로 출력 / PDF 저장")

                st.dataframe(df_code_summary, use_container_width=True)

        # --- Tab 4: 당월 수불 마감 확정 ---
        with tab4:
            st.subheader("🔒 당월 수불 마감 및 전월 DB 확정")
            last_date_str = standard_rows[-1]["date"] if standard_rows else datetime.now().strftime("%Y-%m-%d")
            default_ym = last_date_str[:7]
            closed_ym_input = st.text_input("마감 확정 월 (YYYY-MM)", value=default_ym)

            st.write("▼ **마감 후 다음 달로 이월될 잔여 롯트 미리보기 (실제 기말재고와 100% 동기화)**")
            for item_k, l_list in final_lots.items():
                tot_kg = sum(l['avail_g'] for l in l_list) / 1000.0
                st.write(f"• **{item_k}**: 총 {round(tot_kg, 1)}kg (롯트 {len(l_list)}개 잔여)")
                for l in l_list:
                    st.caption(f"   - 실제 입고일: {l['yymmdd']} | 잔여량: {round(l['avail_g']/1000.0, 1)}kg")

            if st.button("🔒 당월 수불 마감 및 DB 확정 저장", type="primary"):
                save_closed_lots(closed_ym_input, final_lots)
                st.balloons()
                st.success(f"✅ **{closed_ym_input}** 수불 마감이 성공적으로 저장되었습니다!")

        # --- Tab 5: 전처리 공정일지 생성 (빈칸용) ---
        with tab5:
            st.subheader("📋 전처리 공정일지 자동 생성 (수기 작성용)")
            if not df_result.empty:
                unique_dates = sorted(df_result["생산일자"].unique())
                
                col_date, col_empty = st.columns([1, 2])
                with col_date:
                    selected_date = st.selectbox("📅 출력할 생산일자를 선택하세요 (빈칸용)", unique_dates, key="tab5_date")
                
                df_target = df_result[df_result["생산일자"] == selected_date].copy()
                item_priority = {"양상추": 1, "양배추": 2, "적채": 3, "프릴아이스": 4}
                df_target["_item_sort"] = df_target["품목"].map(lambda x: item_priority.get(x, 99))
                df_target = df_target.sort_values(by=["제품명", "_item_sort"])
                
                html_process_log = render_process_log_html(selected_date, df_target, is_completed=False)
                
                col_p_btn, _ = st.columns([1, 3])
                with col_p_btn:
                    render_direct_print_button(html_process_log, "🖨️ 공정일지 바로 출력 / PDF")
                
                st.info("💡 위 [바로 출력 / PDF] 버튼을 누르면 인쇄 대화상자가 즉시 나타납니다. (가로 방향 인쇄 권장)")
                
                components.html(html_process_log, height=800, scrolling=True)
            else:
                st.warning("데이터를 불러오지 못했습니다. 업로드한 파일을 확인해주세요.")

        # --- Tab 6: 전처리 공정일지 생성 (완성본) ---
        with tab6:
            st.subheader("📋 전처리 공정일지 자동 생성 (완성본)")
            if not df_result.empty:
                unique_dates = sorted(df_result["생산일자"].unique())
                
                col_date_6, col_empty_6 = st.columns([1, 2])
                with col_date_6:
                    selected_date_6 = st.selectbox("📅 출력할 생산일자를 선택하세요 (완성본)", unique_dates, key="tab6_date")
                
                df_target_6 = df_result[df_result["생산일자"] == selected_date_6].copy()
                item_priority = {"양상추": 1, "양배추": 2, "적채": 3, "프릴아이스": 4}
                df_target_6["_item_sort"] = df_target_6["품목"].map(lambda x: item_priority.get(x, 99))
                df_target_6 = df_target_6.sort_values(by=["제품명", "_item_sort"])
                
                html_process_log_completed = render_process_log_html(selected_date_6, df_target_6, is_completed=True)
                
                col_p_btn_6, _ = st.columns([1, 3])
                with col_p_btn_6:
                    render_direct_print_button(html_process_log_completed, "🖨️ 완성본 바로 출력 / PDF")
                
                st.info("💡 위 [바로 출력 / PDF] 버튼을 누르면 인쇄 대화상자가 즉시 나타납니다. (가로 방향 인쇄 권장)")
                
                components.html(html_process_log_completed, height=800, scrolling=True)
            else:
                st.warning("데이터를 불러오지 못했습니다. 업로드한 파일을 확인해주세요.")

    except Exception as e:
        st.error(f"처리 중 오류가 발생했습니다: {e}")
        st.code(traceback.format_exc())
else:
    with tab1:
        st.info("👆 상단에서 1. 샐러드 생산일지 엑셀 파일과 2. 원료 매입 수불부를 업로드해 주세요.")
    with tab2:
        st.info("👆 샐러드 생산일지 엑셀 파일을 업로드하시면 자동으로 정제된 수불부가 생성됩니다.")
    with tab3:
        st.info("👆 샐러드 생산일지 엑셀 파일을 업로드하시면 일자별로 자동 생성된 원료코드 표가 출력됩니다.")
    with tab4:
        st.info("👆 샐러드 생산일지 엑셀 파일을 업로드하시면 마감 대상 잔여 롯트를 확인할 수 있습니다.")
    with tab5:
        st.info("👆 엑셀 파일을 업로드하면 전처리 공정일지를 확인하고 출력할 수 있습니다.")
    with tab6:
        st.info("👆 전처리 공정일지 완성본을 확인할 수 있습니다.")
